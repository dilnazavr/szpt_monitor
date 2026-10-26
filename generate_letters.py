#!/usr/bin/env python3
"""
Генерация писем ДЭР (RU/KZ) + приложение.
Адреса можно задать с CLI/сайта или оставить пустыми для ручного заполнения.
В таблице: сумма по договору, сумма по СЗПТ, разница; БИН и название поставщика.
"""
from __future__ import annotations

import argparse
import sys
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Dict, List, Optional

if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from config import (
    ADDRESSEES,
    DEFAULT_THRESHOLD_LETTER,
    EXECUTOR,
    LAW_TEXT_KZ,
    LAW_TEXT_RU,
    LETTERS_DIR,
    PHONE,
    SIGNATORY_KZ,
    SIGNATORY_RU,
)
from etalon_loader import get_etalon_period_hint, load_etalon_by_period
from api_client import GoszakupAPIError, GoszakupClient, is_aktobe_region, normalize_lot
from matcher import process_lots


def parse_args():
    p = argparse.ArgumentParser(description="Генерация писем ДЭР по СЗПТ")
    p.add_argument("--threshold", type=float, default=DEFAULT_THRESHOLD_LETTER)
    p.add_argument("--to", type=str, default="uz,uo")
    p.add_argument("--from-excel", type=str, default=None)
    p.add_argument("--days", type=int, default=365)
    # Адреса получателей (пусто = шаблон с линией для заполнения)
    p.add_argument("--addr-uz-ru", type=str, default="")
    p.add_argument("--addr-uz-kz", type=str, default="")
    p.add_argument("--addr-uo-ru", type=str, default="")
    p.add_argument("--addr-uo-kz", type=str, default="")
    return p.parse_args()


def load_data_from_scan(threshold: float, days: int) -> List[Dict[str, Any]]:
    etalon_by_period = load_etalon_by_period()
    if not etalon_by_period:
        print("Нет эталона — невозможно сформировать письма.")
        sys.exit(1)
    try:
        client = GoszakupClient()
    except GoszakupAPIError as e:
        print(f"API: {e}")
        sys.exit(1)

    date_to = datetime.now().strftime("%Y-%m-%d")
    date_from = (datetime.now() - timedelta(days=days)).strftime("%Y-%m-%d")
    raw = client.search_lots(date_from=date_from, date_to=date_to)
    lots = [normalize_lot(r) for r in raw if is_aktobe_region(r)]
    results, _ = process_lots(lots, etalon_by_period, threshold=threshold)
    print(f"Для писем (>={threshold}%): {len(results)} фактов")
    return results


def load_data_from_excel(path: str, threshold: float) -> List[Dict[str, Any]]:
    df = pd.read_excel(path)
    colmap = {
        "Год": "year", "Дата": "date", "Период эталона": "etalon_period",
        "Товар": "product", "Предмет лота": "lot_name",
        "Цена за ед.": "unit_price", "Эталон Актобе": "etalon",
        "% отклонения": "percent",
        "Сумма": "amount", "Сумма по договору": "amount",
        "Сумма по цене СЗПТ": "amount_szpt", "Разница": "difference",
        "Количество": "qty",
        "Заказчик": "customer", "БИН заказчика": "customer_bin",
        "Поставщик": "supplier_name", "БИН поставщика": "supplier_biin",
        "Номер лота": "lot_number", "Номер объявления": "number_anno",
        "Номер договора": "contract_number", "Статус": "status",
    }
    df = df.rename(columns={k: v for k, v in colmap.items() if k in df.columns})
    if "percent" not in df.columns:
        print("В Excel нет колонки percent / % отклонения")
        sys.exit(1)
    df = df[df["percent"] >= threshold]
    return df.to_dict(orient="records")


def build_appendix(rows: List[Dict[str, Any]], path: Path) -> None:
    data = []
    for i, r in enumerate(rows, 1):
        amount = r.get("amount") or 0
        amount_szpt = r.get("amount_szpt")
        if amount_szpt is None:
            try:
                et = float(r.get("etalon") or 0)
                qty = float(r.get("qty") or 0)
                amount_szpt = round(et * qty, 2) if et and qty else ""
            except (TypeError, ValueError):
                amount_szpt = ""
        diff = r.get("difference")
        if diff is None and amount_szpt != "":
            try:
                diff = round(float(amount) - float(amount_szpt), 2)
            except (TypeError, ValueError):
                diff = ""
        data.append({
            "№": i,
            "Заказчик": r.get("customer") or "",
            "БИН заказчика": r.get("customer_bin") or "",
            "Поставщик": r.get("supplier_name") or "",
            "БИН поставщика": r.get("supplier_biin") or "",
            "Товар": r.get("product") or "",
            "Период эталона": r.get("etalon_period") or "",
            "Договор / лот": (
                r.get("contract_number")
                or r.get("lot_number")
                or r.get("number_anno")
                or ""
            ),
            "Цена ед.": r.get("unit_price") or "",
            "Эталон Stat (Актобе)": r.get("etalon") or "",
            "%": r.get("percent") or "",
            "Сумма по договору": amount,
            "Сумма по цене СЗПТ": amount_szpt,
            "Разница": diff,
            "Рекомендация": "заключить доп. соглашение",
        })
    pd.DataFrame(data).to_excel(path, index=False, engine="openpyxl")
    print(f"Приложение: {path}")


def _add_paragraph(doc: Document, text: str, bold: bool = False,
                   space_after: int = 6, font_size: int = 12):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(font_size)
    run.font.name = "Times New Roman"
    run.bold = bold
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    return p


def _resolve_address(key: str, lang: str, custom: str) -> str:
    """custom непустой — подставить; иначе шаблон с пустой строкой."""
    if custom and custom.strip():
        # если одна строка ФИО — дополняем должностью из шаблона
        base = ADDRESSEES[key][lang].split("\n")
        if len(base) >= 2 and "\n" not in custom:
            return "\n".join(base[:-1] + [custom.strip()])
        return custom.strip()
    return ADDRESSEES[key][lang]


def create_letter_ru(addressee_key: str, rows: List[Dict[str, Any]],
                     period: str, out_path: Path, address_text: str) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)

    for line in address_text.split("\n"):
        _add_paragraph(doc, line, bold=True)

    _add_paragraph(doc, "")
    today = datetime.now().strftime("%d.%m.%Y")
    _add_paragraph(doc, f"№ ______ от {today}")
    _add_paragraph(doc, "")

    _add_paragraph(
        doc,
        "Департамент экономических расследований по Актюбинской области "
        "осуществляет мониторинг цен на социально значимые продовольственные товары (СЗПТ) "
        "в рамках контроля соблюдения законодательства о торговой деятельности.",
        space_after=8,
    )
    n = len(rows)
    _add_paragraph(
        doc,
        f"По результатам мониторинга за период {period} выявлено {n} "
        f"{'факт' if n == 1 else 'факта' if 2 <= n <= 4 else 'фактов'} "
        f"поставок СЗПТ с превышением установленной торговой надбавки.",
        space_after=8,
    )
    _add_paragraph(doc, LAW_TEXT_RU, space_after=8)
    _add_paragraph(
        doc,
        "Анализ показал, что цены, зафиксированные в договорах (лотах) государственных закупок, "
        "существенно превышают средние розничные цены по городу Актобе "
        "за соответствующий период, опубликованные Бюро национальной статистики (stat.gov.kz).",
        space_after=8,
    )
    _add_paragraph(doc, "В связи с изложенным просим:", space_after=4)
    _add_paragraph(
        doc,
        "1. Пересмотреть действующие договоры на поставку СЗПТ и заключить дополнительные "
        "соглашения в соответствии с таблицей-приложением к настоящему письму.",
        space_after=4,
    )
    _add_paragraph(
        doc,
        "2. Обязать подведомственные учреждения при заключении договоров ориентироваться "
        "на средние цены Бюро национальной статистики по городу Актобе за соответствующий период.",
        space_after=4,
    )
    _add_paragraph(
        doc,
        "3. Проинформировать Департамент о принятых мерах в срок не позднее 10 рабочих дней "
        "с даты получения настоящего письма.",
        space_after=8,
    )
    _add_paragraph(doc, "Контроль за исполнением настоящего письма оставляю за собой.", space_after=12)
    for line in SIGNATORY_RU.split("\n"):
        _add_paragraph(doc, line)
    _add_paragraph(doc, "")
    _add_paragraph(doc, f"Исп.: {EXECUTOR}", font_size=10)
    _add_paragraph(doc, f"Тел.: {PHONE}", font_size=10)
    doc.save(out_path)
    print(f"Письмо RU ({addressee_key}): {out_path}")


def create_letter_kz(addressee_key: str, rows: List[Dict[str, Any]],
                     period: str, out_path: Path, address_text: str) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)

    for line in address_text.split("\n"):
        _add_paragraph(doc, line, bold=True)

    _add_paragraph(doc, "")
    today = datetime.now().strftime("%d.%m.%Y")
    _add_paragraph(doc, f"№ ______ {today} ж.")
    _add_paragraph(doc, "")
    _add_paragraph(
        doc,
        "Ақтөбе облысы бойынша экономикалық тергеу департаменті сауда қызметі туралы "
        "заңнаманың сақталуын бақылау шеңберінде әлеуметтік маңызы бар азық-түлік тауарларының "
        "(ӘМАТ) бағасына мониторинг жүргізеді.",
        space_after=8,
    )
    n = len(rows)
    _add_paragraph(
        doc,
        f"{period} кезеңіне жүргізілген мониторинг нәтижелері бойынша сауда үстемесінің "
        f"белгіленген мөлшерінен асатын ӘМАТ жеткізудің {n} фактісі анықталды.",
        space_after=8,
    )
    _add_paragraph(doc, LAW_TEXT_KZ, space_after=8)
    _add_paragraph(
        doc,
        "Талдау көрсеткендей, мемлекеттік сатып алу шарттарында (лоттарда) бекітілген бағалар "
        "Ақтөбе қаласы бойынша тиісті кезеңдегі Ұлттық статистика бюросы (stat.gov.kz) "
        "жариялаған орташа бөлшек сауда бағаларынан айтарлықтай жоғары.",
        space_after=8,
    )
    _add_paragraph(doc, "Жоғарыда айтылғандарға байланысты сұраймыз:", space_after=4)
    _add_paragraph(
        doc,
        "1. ӘМАТ жеткізуге арналған қолданыстағы шарттарды қайта қарап, осы хатқа қосымша "
        "кестеге сәйкес қосымша келісімдер жасасуды.",
        space_after=4,
    )
    _add_paragraph(
        doc,
        "2. Ведомстволық бағынысты мекемелерге шарттар жасасу кезінде Ұлттық статистика бюросының "
        "Ақтөбе қаласы бойынша тиісті кезеңдегі орташа бағаларына бағдарлануды міндеттеуді.",
        space_after=4,
    )
    _add_paragraph(
        doc,
        "3. Осы хатты алған күннен бастап 10 жұмыс күні ішінде қабылданған шаралар туралы "
        "Департаментті хабардар етуді.",
        space_after=8,
    )
    _add_paragraph(doc, "Осы хаттың орындалуын бақылауды өзіме қалдырамын.", space_after=12)
    for line in SIGNATORY_KZ.split("\n"):
        _add_paragraph(doc, line)
    _add_paragraph(doc, "")
    _add_paragraph(doc, f"Орынд.: {EXECUTOR}", font_size=10)
    _add_paragraph(doc, f"Тел.: {PHONE}", font_size=10)
    doc.save(out_path)
    print(f"Письмо KZ ({addressee_key}): {out_path}")


def main():
    args = parse_args()
    targets = [t.strip().lower() for t in args.to.split(",") if t.strip()]
    for t in targets:
        if t not in ADDRESSEES:
            print(f"Неизвестный адресат: {t}")
            sys.exit(1)

    if args.from_excel:
        rows = load_data_from_excel(args.from_excel, args.threshold)
    else:
        rows = load_data_from_scan(args.threshold, args.days)

    if not rows:
        print("Нет данных для писем (фактов >= порога не найдено).")
        rows = []

    period = get_etalon_period_hint()
    LETTERS_DIR.mkdir(parents=True, exist_ok=True)

    build_appendix(rows, LETTERS_DIR / "Приложение_СЗПТ_таблица.xlsx")

    addr_map = {
        "uz": {"ru": args.addr_uz_ru, "kz": args.addr_uz_kz},
        "uo": {"ru": args.addr_uo_ru, "kz": args.addr_uo_kz},
    }
    for key in targets:
        create_letter_ru(
            key, rows, period,
            LETTERS_DIR / f"Письмо_{key.upper()}_СЗПТ_RU.docx",
            _resolve_address(key, "ru", addr_map[key]["ru"]),
        )
        create_letter_kz(
            key, rows, period,
            LETTERS_DIR / f"Письмо_{key.upper()}_СЗПТ_KZ.docx",
            _resolve_address(key, "kz", addr_map[key]["kz"]),
        )

    print("\nВсе файлы в:", LETTERS_DIR)
    print("Готово.")


if __name__ == "__main__":
    main()
